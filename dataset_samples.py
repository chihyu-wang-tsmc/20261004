"""50 個資料集總表 + 每個內部分類各抽一筆真實資料，輸出 dataset_samples.xlsx / dataset_samples.md。

資料一律從各資料集的官方來源（GitHub raw / 官方 S3、GCS）下載到 dataset_samples_cache/，
不讀 ~/safety_benchmarks 那些本機檔，所以在任何機器上跑出來都一樣。
HuggingFace 上才有的資料集在雲端環境連不到，那幾列標「未取得」，在本機補跑時會照樣標。

標籤欄（refuse / coding / model_route / severity）照 deep12.py 實際使用的值，也就是 deep10.DATASETS；
這裡寫死是為了不 import deep10（它會連帶 import langchain、pyarrow 等）。

    python dataset_samples.py            # 下載（已有就跳過）+ 產生兩個檔
"""

import csv
import gzip
import io
import json
import os
import re
import urllib.request
import zipfile
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "dataset_samples_cache")
R = "https://raw.githubusercontent.com"
csv.field_size_limit(10**9)

SOURCES = {
    "advbench_behaviors.csv": f"{R}/llm-attacks/llm-attacks/main/data/advbench/harmful_behaviors.csv",
    "advbench_strings.csv": f"{R}/llm-attacks/llm-attacks/main/data/advbench/harmful_strings.csv",
    "harmbench.csv": f"{R}/centerforaisafety/HarmBench/main/data/behavior_datasets/harmbench_behaviors_text_all.csv",
    "strongreject.csv": f"{R}/alexandrasouly/strongreject/main/strongreject_dataset/strongreject_dataset.csv",
    "sst.csv": f"{R}/bertiev/SimpleSafetyTests/main/SimpleSafetyTests%20-%20test%20cases.csv",
    "openaimod.jsonl.gz": f"{R}/openai/moderation-api-release/main/data/samples-1680.jsonl.gz",
    "jbb_pair.json": f"{R}/JailbreakBench/artifacts/main/attack-artifacts/PAIR/black_box/gpt-4-0125-preview.json",
    "dan.csv": f"{R}/verazuo/jailbreak_llms/main/data/forbidden_question/forbidden_question_set.csv",
    "tensortrust_hijack.jsonl": f"{R}/HumanCompatibleAI/tensor-trust-data/main/benchmarks/hijacking-robustness/v1/hijacking_robustness_dataset.jsonl",
    "tensortrust_extraction.jsonl": f"{R}/HumanCompatibleAI/tensor-trust-data/main/benchmarks/extraction-robustness/v1/extraction_robustness_dataset.jsonl",
    "bipia_text.json": f"{R}/microsoft/BIPIA/main/benchmark/text_attack_test.json",
    "bipia_code.json": f"{R}/microsoft/BIPIA/main/benchmark/code_attack_test.json",
    "injecagent_dh.json": f"{R}/uiuc-kang-lab/InjecAgent/main/data/test_cases_dh_base.json",
    "injecagent_ds.json": f"{R}/uiuc-kang-lab/InjecAgent/main/data/test_cases_ds_base.json",
    "wmdp.zip": "https://cais-wmdp.s3.us-west-1.amazonaws.com/wmdp-mcqs.zip",  # 官方 README 公開的密碼 wmdpmcqs
    "cybermetric500.json": f"{R}/cybermetric/CyberMetric/main/CyberMetric-500-v1.json",
    **{f"cti_{t}.tsv": f"{R}/xashru/cti-bench/main/data/cti-{t}.tsv" for t in ("mcq", "rcm", "vsp", "ate", "taa")},
    **{f"cse_{k}.json": f"{R}/meta-llama/PurpleLlama/main/CybersecurityBenchmarks/datasets/{p}"
       for k, p in {"mitre": "mitre/mitre_benchmark_100_per_category_with_augmentation.json",
                    "interpreter": "interpreter/interpreter.json",
                    "instruct": "instruct/instruct.json",
                    "autocomplete": "autocomplete/autocomplete.json",
                    "frr": "mitre_frr/mitre_frr.json"}.items()},
    "rmcbench.json": f"{R}/qing-yuan233/RMCBench/main/data/json/prompt.json",
    "secbench_mcq.jsonl": f"{R}/secbench-git/SecBench/main/data/MCQs_2730.jsonl",
    "secbench_saq.jsonl": f"{R}/secbench-git/SecBench/main/data/SAQs_270.jsonl",
    "humaneval.jsonl.gz": f"{R}/openai/human-eval/master/data/HumanEval.jsonl.gz",
    "mbpp_sanitized.json": f"{R}/google-research/google-research/master/mbpp/sanitized-mbpp.json",
    "mbjp.jsonl": f"{R}/amazon-science/mxeval/main/data/mbxp/mbjp_release_v1.2.jsonl",
    "securityeval.jsonl": f"{R}/s2e-lab/SecurityEval/main/dataset.jsonl",
    "gsm8k_test.jsonl": f"{R}/openai/grade-school-math/master/grade_school_math/data/test.jsonl",
    # ARC 官方 zip 有 680MB（含語料庫），只解出兩個 test 檔
    "arc.zip": "https://ai2-public-datasets.s3.amazonaws.com/arc/ARC-V1-Feb2018.zip",
}


def fetch():
    os.makedirs(CACHE, exist_ok=True)
    for name, url in SOURCES.items():
        path = os.path.join(CACHE, name)
        if name == "arc.zip" and os.path.exists(os.path.join(CACHE, "ARC-Challenge-Test.jsonl")):
            continue
        if os.path.exists(path):
            continue
        print("下載", name)
        try:
            urllib.request.urlretrieve(url, path)
        except Exception as e:  # 連不到就讓那一列標「未取得」
            print(f"  失敗 {type(e).__name__}: {e}")
            continue
        if name == "arc.zip":
            with zipfile.ZipFile(path) as z:
                for sub in ("ARC-Challenge", "ARC-Easy"):
                    member = f"ARC-V1-Feb2018-2/{sub}/{sub}-Test.jsonl"
                    with z.open(member) as src, open(os.path.join(CACHE, f"{sub}-Test.jsonl"), "wb") as dst:
                        dst.write(src.read())
            os.remove(path)


# ---- 讀檔小工具 ----

def p(name):
    return os.path.join(CACHE, name)


def rows_csv(name, delim=","):
    with open(p(name), encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter=delim))


def rows_jsonl(name):
    opener = gzip.open if name.endswith(".gz") else open
    with opener(p(name), "rt", encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def rows_json(name):
    with open(p(name), encoding="utf-8") as f:
        return json.load(f)


def clip(text, n=280):
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[:n] + "…"


# 給老闆看的表不放涉及未成年人的內容和露骨的性描寫；樣本碰到這些字就跳過，換下一筆
SKIP = re.compile(r"child porn|underage|minor|teen|kid|child sex|pedo"
                  r"|dildo|cum\b|cock|pussy|dick|fuck|orgasm|horny|naked|nude", re.I)


def pick(texts):
    """挑一筆當代表：第一筆長度適中的（太短看不出樣子、太長塞不進表格），固定順序所以每次一樣。"""
    texts = [t for t in texts if isinstance(t, str) and t.strip()]
    texts = [t for t in texts if not SKIP.search(t)] or texts
    for t in texts:
        if 40 <= len(t) <= 400:
            return clip(t)
    return clip(texts[0]) if texts else ""


def group(items, key, text, order=None):
    """依 key 分組，每組回傳 (分類, 題數, 代表樣本)。"""
    g = defaultdict(list)
    for it in items:
        g[key(it)].append(text(it))
    keys = order or sorted(g, key=lambda k: -len(g[k]))
    return [(k, len(g[k]), pick(g[k])) for k in keys if k in g]


# 不列原文的分類：內容涉及未成年人或兒少虐待，給老闆看的表不放原文
REDACT = "（此分類內容涉及未成年人，不列原文）"


def sample_advbench():
    return [("harmful_behaviors", *(lambda r: (len(r), pick([x["goal"] for x in r])))(rows_csv("advbench_behaviors.csv"))),
            ("harmful_strings", *(lambda r: (len(r), pick([x["target"] for x in r])))(rows_csv("advbench_strings.csv")))]


def sample_harmbench():
    return group(rows_csv("harmbench.csv"), lambda x: x["SemanticCategory"], lambda x: x["Behavior"])


def sample_strongreject():
    return group(rows_csv("strongreject.csv"), lambda x: x["category"], lambda x: x["forbidden_prompt"])


def sample_sst():
    out = group(rows_csv("sst.csv"), lambda x: x["harm_area"], lambda x: x["prompts_final"])
    return [(k, n, REDACT if k == "Child safety" else s) for k, n, s in out]


OPENAI_LABELS = {"S": "S 色情", "H": "H 仇恨", "V": "V 暴力", "HR": "HR 騷擾", "SH": "SH 自傷",
                 "S3": "S3 未成年色情", "H2": "H2 仇恨威脅", "V2": "V2 血腥暴力"}


def sample_openaimod(flagged):
    rows = rows_jsonl("openaimod.jsonl.gz")
    if not flagged:
        clean = [r["prompt"] for r in rows if not any(r.get(k) == 1 for k in OPENAI_LABELS)]
        return [("8 個標籤全為 0", len(clean), pick(clean))]
    out = []
    for k, label in OPENAI_LABELS.items():
        hit = [r for r in rows if r.get(k) == 1]
        # 優先挑「只有這一個標籤」的，比較能代表這個分類
        only = [r["prompt"] for r in hit if sum(r.get(j) == 1 for j in OPENAI_LABELS) == 1]
        # S／S3 本身就是色情文字，挑哪一筆都一樣露骨，不列原文
        sample = REDACT if k == "S3" else "（成人色情文字，不列原文）" if k == "S" else pick(only or [r["prompt"] for r in hit])
        out.append((label, len(hit), sample))
    return out


def sample_jbb_harmful():
    return group(rows_json("jbb_pair.json")["jailbreaks"], lambda x: x["category"], lambda x: x["goal"])


def sample_dan():
    return group(rows_csv("dan.csv"), lambda x: x["content_policy_name"], lambda x: x["question"])


def sample_tensortrust():
    h = rows_jsonl("tensortrust_hijack.jsonl")
    e = rows_jsonl("tensortrust_extraction.jsonl")
    return [("hijacking（deep10 用這組）", len(h), pick([x["attack"] for x in h])),
            ("extraction", len(e), pick([x["attack"] for x in e]))]


def sample_bipia():
    text = rows_json("bipia_text.json")
    code = rows_json("bipia_code.json")
    return ([(f"文字攻擊／{k}", len(v), pick(v)) for k, v in text.items()]
            + [(f"程式攻擊／{k}（deep10 未用）", len(v), pick(v)) for k, v in code.items()])


def sample_injecagent():
    out = []
    for name, tag in (("injecagent_dh.json", "直接傷害 dh"), ("injecagent_ds.json", "資料竊取 ds（deep10 未用）")):
        out += [(f"{tag}／{k}", n, s) for k, n, s in
                group(rows_json(name), lambda x: x["Attack Type"], lambda x: x["Attacker Instruction"])]
    return out


def sample_wmdp():
    out = []
    with zipfile.ZipFile(p("wmdp.zip")) as z:
        for sub, label in (("cyber", "Cyber（deep10 用這組）"), ("bio", "Bio（deep10 未用）"), ("chem", "Chem（deep10 未用）")):
            qs = json.loads(z.read(f"wmdp-mcqs/{sub}_questions.json", pwd=b"wmdpmcqs"))
            out.append((label, len(qs), pick([q["question"] for q in qs])))
    return out


def sample_cybermetric():
    q = rows_json("cybermetric500.json")["questions"]
    return [("（檔案沒有分類欄位）500 題版", len(q), pick([x["question"] for x in q]))]


def sample_ctibench():
    out = []
    for t, col, label in (("mcq", "Question", "CTI-MCQ 知識選擇題（deep10 用這組）"),
                          ("rcm", "Description", "CTI-RCM CVE→CWE 根因對應"),
                          ("vsp", "Description", "CTI-VSP CVSS 向量預測"),
                          ("ate", "Description", "CTI-ATE ATT&CK 技術抽取"),
                          ("taa", "Text", "CTI-TAA 威脅行為者歸因")):
        r = rows_csv(f"cti_{t}.tsv", "\t")
        out.append((label, len(r), pick([x[col] for x in r])))
    return out


def sample_cse(kind):
    r = rows_json(f"cse_{kind}.json")
    if kind == "mitre":
        return group(r, lambda x: x["mitre_category"], lambda x: x["base_prompt"])
    if kind == "interpreter":
        return group(r, lambda x: x["attack_type"][0], lambda x: x["mutated_prompt"])
    if kind == "frr":
        return [("is_malicious=false", len(r), pick([x["mutated_prompt"] for x in r]))]
    return group(r, lambda x: x["language"], lambda x: x["test_case_prompt"])


def sample_rmcbench():
    def key(x):
        lvl = f" Level {x['level']}" if x["task"] == "text-to-code generation" else ""
        return f"{x['category']}／{x['task']}{lvl}"
    return group(rows_json("rmcbench.json"), key, lambda x: x["prompt"],
                 order=sorted({key(x) for x in rows_json("rmcbench.json")}))


def sample_secbench(name):
    return group(rows_jsonl(name), lambda x: x["domain"], lambda x: x["question"])


def sample_single(name, field, label="（沒有子分類）", loader=rows_jsonl):
    r = loader(name)
    return [(label, len(r), pick([x[field] for x in r]))]


def sample_securityeval():
    r = rows_jsonl("securityeval.jsonl")
    out = group(r, lambda x: x["ID"].split("_")[0], lambda x: x["Prompt"])
    top = out[:5]
    return top + [(f"其餘 {len(out) - 5} 種 CWE", sum(n for _, n, _ in out[5:]), "（略）")]


def sample_arc():
    out = []
    for sub in ("ARC-Challenge", "ARC-Easy"):
        r = rows_jsonl(f"{sub}-Test.jsonl")
        out.append((sub, len(r), pick([x["question"]["stem"] + " " + " ".join(
            f"({c['label']}) {c['text']}" for c in x["question"]["choices"]) for x in r])))
    return out


def missing(reason):
    return lambda: [("—", None, f"未取得真實資料：{reason}")]


HF = "資料只放在 HuggingFace，這個環境的網路代理擋住 huggingface.co"

# (類別, 名稱, refuse, coding, model_route, 內容說明, 官網, 取樣函式)
DATASETS = [
    ("safety", "AdvBench", "該擋", "否", "推定 fast",
     "2 個子集：harmful_behaviors 520 條、harmful_strings 574 條。沒有分類欄位。deep10 只取 harmful_behaviors",
     "https://github.com/llm-attacks/llm-attacks", sample_advbench),
    ("safety", "HarmBench", "該擋", "否", "推定 fast",
     "文字版 400 條。功能分類 standard 200／contextual 100／copyright 100，語意分類 7 種（下列）",
     "https://github.com/centerforaisafety/HarmBench", sample_harmbench),
    ("safety", "StrongREJECT", "該擋", "否", "推定 fast",
     "313 條、6 類。221 條自寫，其餘取自 DAN、AdvBench、MaliciousInstruct、HarmfulQ、MasterKey 等",
     "https://github.com/alexandrasouly/strongreject", sample_strongreject),
    ("safety", "SimpleSafetyTests", "該擋", "否", "推定 fast",
     "100 條人工撰寫英文提示、5 個傷害領域，每領域再分「資訊查詢」與「指令／行動」兩種問法",
     "https://github.com/bertiev/SimpleSafetyTests", sample_sst),
    ("safety", "BeaverTails-Eval", "該擋", "否", "推定 fast",
     "700 條，14 個傷害類別各 50 條（動物虐待、兒童虐待、政治爭議、歧視、毒品武器、財產犯罪、仇恨言論、錯誤資訊、"
     "非暴力不道德、隱私、自傷、色情、恐怖組織、暴力教唆）",
     "https://huggingface.co/datasets/PKU-Alignment/BeaverTails-Evaluation", missing(HF)),
    ("safety", "Aegis-unsafe", "該擋", "否", "推定 fast",
     "※NVIDIA Aegis 1.0，約 1.1 萬條（提示取自 HH-RLHF），13 個風險類別＋Needs Caution。deep10 取 unsafe",
     "https://huggingface.co/datasets/nvidia/Aegis-AI-Content-Safety-Dataset-1.0", missing(HF)),
    ("safety", "Aegis-safe", "免擋", "否", "推定 fast", "同上，deep10 取 safe",
     "https://huggingface.co/datasets/nvidia/Aegis-AI-Content-Safety-Dataset-1.0", missing(HF)),
    ("safety", "OpenAIMod-flagged", "免擋", "否", "推定 fast",
     "1,680 條內容審查樣本、8 個二元標籤（一條可有多個）。deep10 取任一標籤為 1",
     "https://github.com/openai/moderation-api-release", lambda: sample_openaimod(True)),
    ("safety", "OpenAIMod-clean", "免擋", "否", "推定 fast", "同上，deep10 取 8 個標籤全為 0",
     "https://github.com/openai/moderation-api-release", lambda: sample_openaimod(False)),
    ("jailbreak", "JBB-harmful", "該擋", "否", "推定 fast",
     "100 條有害行為，依 OpenAI 使用政策分 10 類各 10 條（樣本取自官方 artifacts repo 的 goal 欄）",
     "https://github.com/JailbreakBench/jailbreakbench", sample_jbb_harmful),
    ("jailbreak", "JBB-benign", "免擋", "否", "推定 fast", "100 條無害行為，與有害版主題一一對應",
     "https://github.com/JailbreakBench/jailbreakbench", missing(HF)),
    ("jailbreak", "DAN-forbidden", "該擋", "否", "推定 fast", "390 題，13 個禁止情境各 30 題",
     "https://github.com/verazuo/jailbreak_llms", sample_dan),
    ("jailbreak", "JailbreakTrigger", "該擋", "否", "推定 fast",
     "※TrustLLM safety 子任務，約 1,400 條，用 13～14 種越獄手法包裝有害問題",
     "https://github.com/HowieHwong/TrustLLM", missing("資料放在 HuggingFace／Google Drive，這個環境連不到")),
    ("jailbreak", "WildJailbreak-harmful", "該擋", "否", "推定 fast",
     "訓練集 262K，分 vanilla／adversarial × harmful／benign 4 種。deep10 取 eval 的 adversarial_harmful",
     "https://github.com/allenai/wildteaming", missing(HF)),
    ("jailbreak", "WildJailbreak-benign", "免擋", "否", "推定 fast", "同上，取 eval 的 adversarial_benign",
     "https://github.com/allenai/wildteaming", missing(HF)),
    ("jailbreak", "HarmBench-PAIR", "該擋", "否", "推定 fast",
     "HarmBench 行為 + PAIR 攻擊（攻擊者 LLM 迭代改寫），每個目標模型一份 test cases",
     "https://github.com/centerforaisafety/HarmBench", missing("攻擊產生的 test cases 不在 repo 裡，要自己跑攻擊產生")),
    ("jailbreak", "HarmBench-GCG", "該擋", "否", "推定 fast", "同上，GCG：梯度搜尋的對抗後綴",
     "https://github.com/centerforaisafety/HarmBench", missing("同 HarmBench-PAIR")),
    ("jailbreak", "HarmBench-AutoDAN", "該擋", "否", "推定 fast", "同上，AutoDAN：基因演算法演化的角色扮演提示",
     "https://github.com/centerforaisafety/HarmBench", missing("同 HarmBench-PAIR")),
    ("injection", "deepset-PI", "免擋", "否", "推定 fast",
     "※662 條（train 546／test 116），英德混合，label 1 = injection。deep10 取 test 的 label=1",
     "https://huggingface.co/datasets/deepset/prompt-injections", missing(HF)),
    ("injection", "xTRam1-PI", "該擋", "否", "推定 fast",
     "※約 1 萬條合成資料，label 1 = injection，多為索取憑證、個資。deep10 取 test 的 label=1",
     "https://huggingface.co/datasets/xTRam1/safe-guard-prompt-injection", missing(HF)),
    ("injection", "Gandalf-ignore", "免擋", "否", "推定 fast",
     "※Lakera Gandalf 遊戲中「ignore instructions」類提示約 1,000 條。deep10 取 test",
     "https://huggingface.co/datasets/Lakera/gandalf_ignore_instructions", missing(HF)),
    ("injection", "TensorTrust-hijack", "免擋", "否", "推定 fast",
     "Tensor Trust 線上攻防遊戲收集，benchmark v1 分 hijacking 與 extraction 兩個任務",
     "https://github.com/HumanCompatibleAI/tensor-trust-data", sample_tensortrust),
    ("injection", "BIPIA-attacks", "免擋", "否", "推定 fast",
     "間接注入，5 個應用情境（email／web／table QA、摘要、code QA）。deep10 只取文字攻擊 75 條",
     "https://github.com/microsoft/BIPIA", sample_bipia),
    ("injection", "InjecAgent-dh", "免擋", "否", "推定 fast",
     "工具代理間接注入。dh 510 例（30 種攻擊指令 × 17 種使用者工具），deep10 只取 dh 的 Attacker Instruction",
     "https://github.com/uiuc-kang-lab/InjecAgent", sample_injecagent),
    ("cyber", "WMDP-Cyber", "免擋", "否", "推定 fast",
     "WMDP 共 3,668 題危險知識選擇題，分 Bio／Cyber／Chem 3 個子集",
     "https://github.com/centerforaisafety/wmdp", sample_wmdp),
    ("cyber", "CyberMetric", "免擋", "否", "推定 fast",
     "80／500／2,000／10,000 題 4 個版本，RAG 生成、專家驗證的 4 選 1 題。deep10 用 500 題版",
     "https://github.com/cybermetric/CyberMetric", sample_cybermetric),
    ("cyber", "SecEval", "免擋", "否", "推定 fast",
     "2,126 題，9 個主題：SystemSecurity 1065、ApplicationSecurity 808、WebSecurity 773、PenTest 475、"
     "Vulnerability 334、SoftwareSecurity 232、NetworkSecurity 230、MemorySafety 48、Cryptography 14",
     "https://github.com/XuanwuAI/SecEval", missing("題目檔只放在 HuggingFace")),
    ("cyber", "CTIBench-mcq", "免擋", "否", "推定 fast", "威脅情報 benchmark，5 個任務。deep10 只取 MCQ",
     "https://github.com/xashru/cti-bench", sample_ctibench),
    ("cyber", "CyberSecEval-mitre", "該擋", "是", "推定 fast", "1,000 條，MITRE ATT&CK 10 個戰術各 100。deep10 取 base_prompt",
     "https://github.com/meta-llama/PurpleLlama/tree/main/CybersecurityBenchmarks", lambda: sample_cse("mitre")),
    ("cyber", "CyberSecEval-interpreter", "該擋", "是", "推定 fast", "500 條，誘導在程式直譯器裡做壞事，5 類各 100",
     "https://github.com/meta-llama/PurpleLlama/tree/main/CybersecurityBenchmarks", lambda: sample_cse("interpreter")),
    ("cyber", "RMCBench-malcode", "該擋", "是", "推定 fast",
     "惡意程式生成 473 條：text-to-code（Level 1～3，Level 3 套越獄模板）與 code-to-code（補全、翻譯）",
     "https://github.com/qing-yuan233/RMCBench", sample_rmcbench),
    ("cyber", "mitre_frr-benign", "免擋", "是", "推定 fast", "CyberSecEval False Refusal Rate 組：看似攻擊、實為正當的資安程式請求",
     "https://github.com/meta-llama/PurpleLlama/tree/main/CybersecurityBenchmarks", lambda: sample_cse("frr")),
    ("cyber", "SecBench-mcq", "免擋", "否", "推定 fast", "2,730 題（中 2,069／英 661），能力分知識記憶／邏輯推理，領域 9 個",
     "https://github.com/secbench-git/SecBench", lambda: sample_secbench("secbench_mcq.jsonl")),
    ("cyber", "SecBench-saq", "免擋", "否", "推定 fast", "簡答題版 270 題，同樣 9 個領域各 30",
     "https://github.com/secbench-git/SecBench", lambda: sample_secbench("secbench_saq.jsonl")),
    ("coding", "HumanEval", "免擋", "是", "推定 powerful", "164 題手寫 Python 函式補全（簽名 + docstring），單元測試算 pass@k",
     "https://github.com/openai/human-eval", lambda: sample_single("humaneval.jsonl.gz", "prompt")),
    ("coding", "MBPP", "免擋", "是", "推定 powerful", "原版約 974 題入門 Python 題，sanitized 版 427 題。deep10 用 sanitized test",
     "https://github.com/google-research/google-research/tree/master/mbpp",
     lambda: sample_single("mbpp_sanitized.json", "prompt", "sanitized（全部）", rows_json)),
    ("coding", "MBXP-java", "免擋", "是", "推定 powerful", "MBPP 轉成 10 多種語言，deep10 只取 Java 版 mbjp",
     "https://github.com/amazon-science/mxeval", lambda: sample_single("mbjp.jsonl", "prompt", "Java（mbjp）")),
    ("coding", "APPS", "免擋", "是", "推定 powerful",
     "1 萬題（train／test 各 5,000），test 依難度分 Introductory 1,000／Interview 3,000／Competition 1,000",
     "https://github.com/hendrycks/apps", missing("資料檔在 HuggingFace 和 berkeley.edu，這個環境都連不到")),
    ("coding", "CodeContests", "免擋", "是", "推定 powerful",
     "AlphaCode 競賽題，來源 Aizu／AtCoder／CodeChef／Codeforces／HackerEarth，test 165 題",
     "https://github.com/google-deepmind/code_contests", missing("官方檔是 riegeli 格式，這個環境裝不了 riegeli 解析套件")),
    ("coding", "BigCodeBench", "免擋", "是", "推定 powerful",
     "1,140 題實務 Python 任務（139 個函式庫），Complete／Instruct 兩個 split。deep10 取 instruct_prompt",
     "https://github.com/bigcode-project/bigcodebench", missing(HF)),
    ("coding", "LiveCodeBench", "免擋", "是", "推定 powerful",
     "持續從 LeetCode／AtCoder／Codeforces 收新題，release_v1～v6 累積 400→1,055 題。deep10 取 test6",
     "https://github.com/LiveCodeBench/LiveCodeBench", missing(HF)),
    ("coding", "MathQA-Python", "免擋", "否", "推定 powerful", "※MathQA 數學文字題配 Python 解答，約 2.4 萬題",
     "https://github.com/google-research/google-research", missing(HF)),
    ("coding", "SecurityEval", "免擋", "是", "推定 powerful", "121 條程式生成提示、69 種 CWE，每題附一份不安全範例程式",
     "https://github.com/s2e-lab/SecurityEval", sample_securityeval),
    ("coding", "CyberSecEval-instruct", "免擋", "是", "推定 powerful", "1,916 條取自真實開源程式，8 種語言、50 種 CWE",
     "https://github.com/meta-llama/PurpleLlama/tree/main/CybersecurityBenchmarks", lambda: sample_cse("instruct")),
    ("coding", "CyberSecEval-autocomplete", "免擋", "是", "推定 powerful", "同上 1,916 條，改成「給前文、續寫程式」",
     "https://github.com/meta-llama/PurpleLlama/tree/main/CybersecurityBenchmarks", lambda: sample_cse("autocomplete")),
    ("skill", "skill-queries", "免擋", "否", "推定 powerful", "內部資料：要 agent 產出 xlsx、投影片、文件的正常請求",
     "無（內部 ~/jevk5/skill_queries/all.jsonl）", missing("內部資料，不在這個 repo")),
    ("fast-powerful", "GPQA", "免擋", "否", "實測", "研究所程度 4 選 1，生物／物理／化學，Extended 546／Main 448／Diamond 198，deep10 用 Diamond",
     "https://github.com/idavidrein/gpqa", missing("資料集條款要求不公開題目，故不列原文")),
    ("fast-powerful", "AI2ARC", "免擋", "否", "實測", "小學自然科選擇題，ARC-Challenge 與 ARC-Easy 兩個子集",
     "https://allenai.org/data/arc", sample_arc),
    ("fast-powerful", "AIME", "免擋", "否", "實測", "美國數學邀請賽，每年 I、II 兩場各 15 題，答案為 0～999 整數。repo 用 2025、2026 共 60 題",
     "https://artofproblemsolving.com/wiki/index.php/AIME_Problems_and_Solutions",
     missing("沒有官方資料檔，AoPS 在這個環境連不到")),
    ("fast-powerful", "GSM8K", "免擋", "否", "實測", "8.5K 題小學數學文字題（train 7.5K／test 1,319），每題 2～8 步",
     "https://github.com/openai/grade-school-math", lambda: sample_single("gsm8k_test.jsonl", "question", "test")),
]

HEADER = ["類別", "資料集", "refuse GT", "coding GT", "model_route GT", "severity GT",
          "資料集內容說明", "內部分類／場景", "該分類題數", "真實資料舉例", "官網"]


def build():
    out = []  # (資料集層級欄位, [(分類, 題數, 樣本)])
    for cat, name, ref, cod, route, desc, url, fn in DATASETS:
        try:
            subs = fn()
        except (OSError, KeyError, ValueError, zipfile.BadZipFile) as e:
            subs = [("—", None, f"未取得真實資料：讀檔失敗 {type(e).__name__}")]
        out.append(((cat, name, ref, cod, route, "無", desc, url), subs))
    return out


def write_xlsx(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "50 個資料集總表"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    fills = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
    row = 2
    for i, ((cat, name, ref, cod, route, sev, desc, url), subs) in enumerate(data):
        start = row
        for sub, n, sample in subs:
            ws.append([cat, name, ref, cod, route, sev, desc, sub, n, sample, url])
            row += 1
        for r in range(start, row):
            for c in ws[r]:
                c.fill = fills[i % 2]
                c.alignment = Alignment(vertical="top", wrap_text=True)
        if row - start > 1:  # 資料集層級的欄位合併成一格，分類那三欄逐列展開
            for col in (1, 2, 3, 4, 5, 6, 7, 11):
                ws.merge_cells(start_row=start, end_row=row - 1, start_column=col, end_column=col)
    for col, w in zip("ABCDEFGHIJK", (12, 22, 9, 9, 13, 9, 40, 30, 9, 70, 30)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def write_md(data, path):
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")
    lines = ["# 50 個資料集總表（含每個內部分類的真實資料）", "",
             "由 `dataset_samples.py` 產生。標籤欄照 deep12.py 實際使用的 deep10.DATASETS；"
             "「真實資料舉例」是從官方來源下載的原始檔，每個分類挑第一筆長度適中的題目（超過 280 字截斷）。"
             "※ 表示該說明無法從官方頁面核對（HuggingFace／arXiv 在產生環境被擋）。", "",
             "| " + " | ".join(HEADER) + " |", "|" + "---|" * len(HEADER)]
    for (cat, name, ref, cod, route, sev, desc, url), subs in data:
        for j, (sub, n, sample) in enumerate(subs):
            head = [cat, f"**{name}**", ref, cod, route, sev, desc] if j == 0 else [""] * 7
            lines.append("| " + " | ".join(esc(x) for x in head + [sub, "" if n is None else n, sample, url if j == 0 else ""]) + " |")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


if __name__ == "__main__":
    fetch()
    data = build()
    write_xlsx(data, os.path.join(HERE, "dataset_samples.xlsx"))
    write_md(data, os.path.join(HERE, "dataset_samples.md"))
    got = sum(1 for _, subs in data if subs[0][1] is not None)
    print(f"完成：{len(data)} 個資料集，{got} 個有真實資料，共 {sum(len(s) for _, s in data)} 列")
